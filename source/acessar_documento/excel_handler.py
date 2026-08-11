import os
import logging
from pathlib import Path
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

class ExcelHandler:
    """Gerenciador de criação, formatação e atualização da Planilha Mestra Excel (openpyxl)."""

    HEADERS = [
        "ID",
        "Data Processamento",
        "Nome",
        "CPF",
        "Data de Nascimento",
        "Endereço",
        "Arquivo Origem",
        "Status Extração"
    ]

    def __init__(self, file_path: Optional[Path] = None):
        if file_path:
            self.file_path = Path(file_path)
        else:
            base_dir = Path(__file__).resolve().parent.parent.parent
            env_path = os.getenv("PLANILHA_MESTRA_PATH", "output/planilha_mestra.xlsx")
            self.file_path = base_dir / env_path

        self.ensure_file_exists()

    def ensure_file_exists(self) -> None:
        """Cria o diretório pai e a planilha Excel com cabeçalhos se o arquivo não existir."""
        self.file_path.parent.mkdir(parents=True, exist_ok=True)

        if not self.file_path.exists():
            logger.info(f"Criando nova Planilha Mestra em: {self.file_path}")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Dados Extraídos"

            # Escreve cabeçalhos
            ws.append(self.HEADERS)
            self._apply_header_styles(ws)

            wb.save(self.file_path)
            wb.close()

    def _apply_header_styles(self, ws) -> None:
        """Aplica estilo visual profissional aos cabeçalhos (Azul escuro, negrito, centralizado)."""
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_num, _ in enumerate(self.HEADERS, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        ws.row_dimensions[1].height = 25

    def get_next_id(self, ws) -> int:
        """Calcula o próximo ID sequencial na planilha."""
        max_row = ws.max_row
        if max_row <= 1:
            return 1
        last_id_val = ws.cell(row=max_row, column=1).value
        try:
            return int(last_id_val) + 1
        except (ValueError, TypeError):
            return max_row

    def append_record(self, data: Dict[str, str]) -> bool:
        """
        Adiciona um novo registro de dados extraídos na Planilha Mestra.

        :param data: Dicionário contendo nome, cpf, data_nascimento, endereco, status, arquivo_origem.
        :return: True se salvo com sucesso, False caso contrário.
        """
        try:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb.active

            next_id = self.get_next_id(ws)
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            row_data = [
                next_id,
                timestamp,
                data.get("nome", "NÃO ENCONTRADO"),
                data.get("cpf", "NÃO ENCONTRADO"),
                data.get("data_nascimento", "NÃO ENCONTRADO"),
                data.get("endereco", "NÃO ENCONTRADO"),
                data.get("arquivo_origem", ""),
                data.get("status", "PROCESSADO")
            ]

            ws.append(row_data)
            current_row = ws.max_row

            # Formatação simples das células da nova linha
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.border = thin_border
                if col_idx in (1, 2, 4, 5, 8):  # Centralizar ID, Data, CPF, DtNasc, Status
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            self._adjust_column_widths(ws)

            wb.save(self.file_path)
            wb.close()
            logger.info(f"Registro ID {next_id} salvo com sucesso na Planilha Mestra.")
            return True

        except Exception as e:
            logger.error(f"Erro ao salvar registro na Planilha Mestra '{self.file_path}': {e}", exc_info=True)
            return False

    def _adjust_column_widths(self, ws) -> None:
        """Ajusta a largura das colunas automaticamente conforme o conteúdo."""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
