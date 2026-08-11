import logging
from pathlib import Path
from typing import Optional

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

logger = logging.getLogger(__name__)

class DocumentReader:
    """Leitor de conteúdo de documentos (PDF, DOCX, TXT/CSV)."""

    @classmethod
    def read_text(cls, file_path: Path) -> str:
        """
        Lê o conteúdo em texto do arquivo informado com base na sua extensão.
        
        :param file_path: Caminho do arquivo a ser lido.
        :return: Texto extraído do arquivo ou string vazia se falhar.
        """
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            logger.error(f"Arquivo não encontrado ou inválido: {file_path}")
            return ""

        suffix = path.suffix.lower()

        try:
            if suffix == ".pdf":
                return cls._read_pdf(path)
            elif suffix in (".docx", ".doc"):
                return cls._read_docx(path)
            elif suffix in (".txt", ".csv", ".log"):
                return cls._read_plain_text(path)
            else:
                logger.warning(f"Extensão de arquivo não suportada para leitura direta de texto: {suffix}")
                return cls._read_plain_text(path)
        except Exception as e:
            logger.error(f"Erro ao ler o arquivo '{file_path}': {e}", exc_info=True)
            return ""

    @staticmethod
    def _read_pdf(path: Path) -> str:
        """Lê o texto de todas as páginas de um PDF usando pypdf."""
        if pypdf is None:
            logger.error("Biblioteca pypdf não instalada.")
            return ""

        text_content = []
        with open(path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    text_content.append(extracted)
        
        return "\n".join(text_content)

    @staticmethod
    def _read_docx(path: Path) -> str:
        """Lê o texto de parágrafos e tabelas de um documento DOCX usando python-docx."""
        if docx is None:
            logger.error("Biblioteca python-docx não instalada.")
            return ""

        doc = docx.Document(path)
        full_text = []

        # Lê parágrafos
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)

        # Lê células de tabelas se houver
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))

        return "\n".join(full_text)

    @staticmethod
    def _read_plain_text(path: Path) -> str:
        """Lê texto puro de arquivo TXT/CSV tentando UTF-8 e Latin-1."""
        for encoding in ["utf-8", "latin-1", "iso-8859-1"]:
            try:
                with open(path, "r", encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        logger.error(f"Não foi possível decodificar o arquivo de texto '{path}'.")
        return ""
