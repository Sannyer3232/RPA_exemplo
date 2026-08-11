import sys
from pathlib import Path

# Adiciona a raiz do projeto ao sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pytest
from unittest.mock import MagicMock, patch, mock_open
from tempfile import TemporaryDirectory
import openpyxl

from source.acessar_documento.document_reader import DocumentReader
from source.acessar_documento.data_extractor import DataExtractor
from source.acessar_documento.excel_handler import ExcelHandler
from source.acessar_documento.process_documents import DocumentProcessor


# ==========================================
# Testes de Leitura de Documentos (DocumentReader)
# ==========================================
def test_document_reader_non_existent_file():
    text = DocumentReader.read_text(Path("arquivo_inexistente_123.pdf"))
    assert text == ""


def test_document_reader_plain_text():
    with TemporaryDirectory() as temp_dir:
        file_path = Path(temp_dir) / "teste.txt"
        file_path.write_text("Nome: Carlos Oliveira\nCPF: 123.456.789-00", encoding="utf-8")
        
        text = DocumentReader.read_text(file_path)
        assert "Carlos Oliveira" in text
        assert "123.456.789-00" in text


@patch("source.acessar_documento.document_reader.pypdf.PdfReader")
def test_document_reader_pdf(mock_pdf_reader):
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Nome: Ana Souza\nCPF: 987.654.321-11"
    mock_reader_instance = MagicMock()
    mock_reader_instance.pages = [mock_page]
    mock_pdf_reader.return_value = mock_reader_instance

    with TemporaryDirectory() as temp_dir:
        pdf_path = Path(temp_dir) / "documento.pdf"
        pdf_path.write_bytes(b"%PDF-1.4 mock content")

        text = DocumentReader.read_text(pdf_path)
        assert "Ana Souza" in text
        assert "987.654.321-11" in text


@patch("source.acessar_documento.document_reader.docx.Document")
def test_document_reader_docx(mock_docx_doc):
    mock_para = MagicMock()
    mock_para.text = "Nome: Roberto Santos"
    
    mock_cell = MagicMock()
    mock_cell.text = "CPF: 111.222.333-44"
    mock_row = MagicMock()
    mock_row.cells = [mock_cell]
    mock_table = MagicMock()
    mock_table.rows = [mock_row]

    mock_doc_instance = MagicMock()
    mock_doc_instance.paragraphs = [mock_para]
    mock_doc_instance.tables = [mock_table]
    mock_docx_doc.return_value = mock_doc_instance

    with TemporaryDirectory() as temp_dir:
        docx_path = Path(temp_dir) / "contrato.docx"
        docx_path.write_bytes(b"PK mock docx zip")

        text = DocumentReader.read_text(docx_path)
        assert "Roberto Santos" in text
        assert "111.222.333-44" in text


# ==========================================
# Testes do DataExtractor (Regex & Regras)
# ==========================================
def test_data_extractor_cpf():
    text_formatted = "O cliente com CPF 123.456.789-00 solicitou atendimento."
    assert DataExtractor.extract_cpf(text_formatted) == "123.456.789-00"

    text_unformatted = "CPF do titular: 98765432199 no cadastro."
    assert DataExtractor.extract_cpf(text_unformatted) == "987.654.321-99"

    text_invalid = "CPF invalido: 00000000000"
    assert DataExtractor.extract_cpf(text_invalid) == "NÃO ENCONTRADO"


def test_data_extractor_dob():
    text_with_label = "Data de Nascimento: 15/08/1992 na certidão."
    assert DataExtractor.extract_dob(text_with_label) == "15/08/1992"

    text_unlabeled = "Nascido em 25-12-1980 em São Paulo."
    assert DataExtractor.extract_dob(text_unlabeled) == "25-12-1980"

    text_no_date = "Nenhuma data presente neste documento."
    assert DataExtractor.extract_dob(text_no_date) == "NÃO ENCONTRADO"


def test_data_extractor_name():
    text_label = "Nome do Cliente: Juliana de Alencar\nCPF: 123.456.789-00"
    assert DataExtractor.extract_name(text_label) == "Juliana De Alencar"

    text_fallback = "Marcos Vinicius Costa\nRua das Palmeiras, 100"
    assert DataExtractor.extract_name(text_fallback) == "Marcos Vinicius Costa"

    text_no_name = "1234567\n000000\n///"
    assert DataExtractor.extract_name(text_no_name) == "NÃO ENCONTRADO"


def test_data_extractor_address():
    text_label = "Endereço: Avenida Paulista, 1000 - Bela Vista - São Paulo / SP"
    assert DataExtractor.extract_address(text_label) == "Avenida Paulista, 1000 - Bela Vista - São Paulo / SP"

    text_street = "Residente na Rua das Flores 450, Centro"
    assert DataExtractor.extract_address(text_street) == "Rua das Flores 450, Centro"


def test_data_extractor_extract_all_complete():
    doc_text = (
        "CADASTRO DE CLIENTE\n"
        "Nome: Fernanda Lima\n"
        "CPF: 555.444.333-22\n"
        "Data de Nascimento: 10/10/1995\n"
        "Endereço: Alameda Santos 800 - Cerqueira César\n"
    )
    result = DataExtractor.extract_all(doc_text, source_filename="ficha_fernanda.pdf")

    assert result["nome"] == "Fernanda Lima"
    assert result["cpf"] == "555.444.333-22"
    assert result["data_nascimento"] == "10/10/1995"
    assert result["endereco"] == "Alameda Santos 800 - Cerqueira César"
    assert result["status"] == "COMPLETO"
    assert result["arquivo_origem"] == "ficha_fernanda.pdf"


def test_data_extractor_extract_all_empty():
    result = DataExtractor.extract_all("   ", source_filename="vazio.pdf")
    assert result["status"] == "ERRO_LEITURA_VAZIA"


# ==========================================
# Testes do ExcelHandler (openpyxl)
# ==========================================
def test_excel_handler_creation_and_append():
    with TemporaryDirectory() as temp_dir:
        excel_path = Path(temp_dir) / "planilha_mestra_teste.xlsx"
        handler = ExcelHandler(file_path=excel_path)

        assert excel_path.exists()

        # Adiciona primeiro registro
        record1 = {
            "nome": "Pedro Alves",
            "cpf": "111.222.333-44",
            "data_nascimento": "01/01/1990",
            "endereco": "Rua Um, 10",
            "status": "COMPLETO",
            "arquivo_origem": "doc1.pdf"
        }
        success1 = handler.append_record(record1)
        assert success1 is True

        # Adiciona segundo registro
        record2 = {
            "nome": "Carla Dias",
            "cpf": "999.888.777-66",
            "data_nascimento": "02/02/1992",
            "endereco": "Rua Dois, 20",
            "status": "PARCIAL",
            "arquivo_origem": "doc2.txt"
        }
        success2 = handler.append_record(record2)
        assert success2 is True

        # Abre a planilha com openpyxl e valida dados
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        assert ws.max_row == 3  # Cabeçalho + 2 registros

        # Linha 1: Cabeçalhos
        assert ws.cell(row=1, column=1).value == "ID"
        assert ws.cell(row=1, column=3).value == "Nome"

        # Linha 2: Registro 1 (ID 1)
        assert ws.cell(row=2, column=1).value == 1
        assert ws.cell(row=2, column=3).value == "Pedro Alves"
        assert ws.cell(row=2, column=4).value == "111.222.333-44"

        # Linha 3: Registro 2 (ID 2)
        assert ws.cell(row=3, column=1).value == 2
        assert ws.cell(row=3, column=3).value == "Carla Dias"
        wb.close()


def test_excel_handler_exception():
    handler = ExcelHandler()
    with patch("openpyxl.load_workbook", side_effect=Exception("Erro ao abrir planilha")):
        result = handler.append_record({"nome": "Teste"})
        assert result is False


# ==========================================
# Testes do DocumentProcessor (Orquestrador)
# ==========================================
def test_document_processor_flow():
    with TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        input_dir = temp_path / "Documentos_OK"
        input_dir.mkdir()

        # Cria 2 arquivos de teste
        doc1 = input_dir / "cliente1.txt"
        doc1.write_text(
            "Nome: Lucas Mendes\n"
            "CPF: 123.123.123-12\n"
            "Data de Nascimento: 12/12/1988\n"
            "Endereço: Rua Central, 50\n"
        )

        doc2 = input_dir / "cliente2.txt"
        doc2.write_text("Arquivo sem dados extraiveis")

        excel_path = temp_path / "planilha_output.xlsx"
        processor = DocumentProcessor(input_dir=input_dir, excel_path=excel_path)

        stats = processor.process_all_documents()
        assert stats["documentos_encontrados"] == 2
        assert stats["documentos_processados"] == 2
        assert stats["registros_completos"] == 1

        assert excel_path.exists()


def test_document_processor_empty_directory():
    with TemporaryDirectory() as temp_dir:
        empty_dir = Path(temp_dir) / "PastaVazia"
        processor = DocumentProcessor(input_dir=empty_dir)
        stats = processor.process_all_documents()
        assert stats["documentos_encontrados"] == 0


if __name__ == "__main__":
    pytest.main(["-v", __file__])
